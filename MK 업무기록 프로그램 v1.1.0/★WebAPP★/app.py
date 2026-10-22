# -*- coding: utf-8 -*-
"""
MK 업무기록 프로그램 - 웹(Flask) 버전
데스크톱(tkinter) 버전의 핵심 기능을 웹에서 여러 사용자가 동시에 쓸 수 있도록 이식.
"""
import io
import json
import os
import re
import threading
import uuid
from datetime import datetime

from flask import Flask, jsonify, render_template, request, send_file, session

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)

app = Flask(__name__)

# ------------------------------------------------------------------
# 세션용 secret key (재시작해도 로그인 유지되도록 파일에 저장)
# ------------------------------------------------------------------
SECRET_KEY_FILE = os.path.join(DATA_DIR, ".secret_key")
if os.path.exists(SECRET_KEY_FILE):
    with open(SECRET_KEY_FILE, "rb") as f:
        app.secret_key = f.read()
else:
    key = os.urandom(32)
    with open(SECRET_KEY_FILE, "wb") as f:
        f.write(key)
    app.secret_key = key

app.config["SESSION_COOKIE_NAME"] = "mk_worklog_session"
app.config["PERMANENT_SESSION_LIFETIME"] = 60 * 60 * 24 * 365  # 1년

# ------------------------------------------------------------------
# 데스크톱 버전과 동일한 정적 데이터 (작업코드 / 공정명 / 매핑)
# ------------------------------------------------------------------
CODES = ["CAD", "KT", "ET", "KC", "TT"]

CATEGORY_LIST = [
    # CAD 작업코드별 공정명
    "선행검토", "구조 설계", "구조 DR", "내부 DR", "금형 DR", "M/B 발주",
    "건드릴/황삭 DATA", "3D 설계", "수정 데이터 반영", "해석데이터 발송",
    "2D 설계", "자재표", "자료발주", "도면출도",

    # KT 작업코드별 공정명
    "금형 기본자료 작성", "제품 검토서 작성", "수정 데이터 점검/검토",
    "설계 검토 및 지원", "현장대응", "시사출 보고서 작성",
    "설계 외주관리", "설계 데이터 발송", "팩터 DATA 발송", "승인원 작성",
    "인자표 작성", "PSR 작성", "스페어리스트 견적/승인/발주",
    "코팅리스트 견적/승인/발주", "출장업무(공무관련, 이동시간)",
    "사출지원", "기타서류 업무",

    # ET 작업코드별 공정명
    "TO 이후 수정", "제품 데이터 수정 검토", "수정검토서 작성",
    "설계변경 도면작성 및 출도",

    # KC 작업코드별 공정명
    "수주 전 견적기초도면 작성", "수주 전 제품데이터 검토", "일반회의",
    "교육(업무관련)", "표준화 작업", "심사관련 업무", "기타서류 업무",

    # TT 작업코드별 공정명
    "교육(업무무관)", "개인업무능력 STUDY", "외출", "청소", "무(일 없는 경우)",
]

CATEGORY_CODE_MAP = {
    "선행검토": "CAD", "구조 설계": "CAD", "구조 DR": "CAD", "내부 DR": "CAD",
    "금형 DR": "CAD", "M/B 발주": "CAD", "건드릴/황삭 DATA": "CAD",
    "3D 설계": "CAD", "수정 데이터 반영": "CAD", "해석데이터 발송": "CAD",
    "2D 설계": "CAD", "자재표": "CAD", "자료발주": "CAD", "도면출도": "CAD",

    "금형 기본자료 작성": "KT", "제품 검토서 작성": "KT", "수정 데이터 점검/검토": "KT",
    "설계 검토 및 지원": "KT", "현장대응": "KT", "시사출 보고서 작성": "KT",
    "설계 외주관리": "KT", "설계 데이터 발송": "KT", "팩터 DATA 발송": "KT",
    "승인원 작성": "KT", "인자표 작성": "KT", "PSR 작성": "KT",
    "스페어리스트 견적/승인/발주": "KT", "코팅리스트 견적/승인/발주": "KT",
    "출장업무(공무관련, 이동시간)": "KT", "사출지원": "KT",

    "TO 이후 수정": "ET", "제품 데이터 수정 검토": "ET", "수정검토서 작성": "ET",
    "설계변경 도면작성 및 출도": "ET",

    "수주 전 견적기초도면 작성": "KC", "수주 전 제품데이터 검토": "KC", "일반회의": "KC",
    "교육(업무관련)": "KC", "표준화 작업": "KC", "심사관련 업무": "KC", "기타서류 업무": "KC",

    "교육(업무무관)": "TT", "개인업무능력 STUDY": "TT", "외출": "TT", "청소": "TT",
    "무(일 없는 경우)": "TT",
}

MK_PATTERN = re.compile(r"^\d{5}[MTRmtr]$")
MK_TYPING_PATTERN = re.compile(r"^(\d{5})([mtr])$", re.IGNORECASE)
MK_SUFFIX_SOURCE_PATTERN = re.compile(r"^\d{5}[MTRmtr]?$")

TODO_SORTABLE_COLUMNS = {"created_date", "deadline", "dday"}

# 진행 중인 업무(사용자별) - 메모리에만 보관 (데스크톱 버전도 재시작하면 사라짐)
ACTIVE_LOCK = threading.Lock()
ACTIVE_WORK = {}


# ------------------------------------------------------------------
# 사용자별 데이터 저장 (이름만으로 구분, 비밀번호 없음)
# ------------------------------------------------------------------
def slugify_user(name):
    name = (name or "").strip()[:20]
    slug = re.sub(r"[^0-9A-Za-z가-힣_\- ]", "_", name).strip()
    slug = slug.replace(" ", "_")
    return slug or "user"


def user_dir(slug):
    d = os.path.join(DATA_DIR, slug)
    os.makedirs(d, exist_ok=True)
    return d


def settings_path(slug):
    return os.path.join(user_dir(slug), "settings.json")


def logs_path(slug):
    return os.path.join(user_dir(slug), "logs.json")


def load_settings(slug):
    p = settings_path(slug)
    data = {}
    if os.path.exists(p):
        try:
            with open(p, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            data = {}
    data.setdefault("favorites", {})
    data.setdefault("todos", {})
    return data


def save_settings(slug, data):
    with open(settings_path(slug), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_logs(slug):
    p = logs_path(slug)
    if os.path.exists(p):
        try:
            with open(p, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []


def save_logs(slug, logs):
    with open(logs_path(slug), "w", encoding="utf-8") as f:
        json.dump(logs, f, ensure_ascii=False, indent=2)


def current_user_slug():
    return session.get("user_slug")


def login_required_response():
    return jsonify({"error": "login_required", "message": "이름을 먼저 입력해주세요."}), 401


# ------------------------------------------------------------------
# 업무 코드 / MK번호 검증 (데스크톱 버전과 동일한 규칙)
# ------------------------------------------------------------------
def validate_mk(mk_number):
    if not mk_number:
        return True, ""
    if not MK_PATTERN.match(mk_number):
        return False, "MK번호는 5자리 숫자 + 알파벳(M/T/R)이어야 합니다.\n예: 25001M, 25002T, 25003R"
    return True, ""


def normalize_mk(mk_number):
    if not mk_number:
        return mk_number
    m = MK_TYPING_PATTERN.match(mk_number)
    if m:
        return m.group(1) + m.group(2).upper()
    return mk_number


def validate_start_fields(code, mk_number, category):
    if not code:
        return "작업코드를 선택해주세요."
    if code in ("KC", "TT"):
        if not category:
            return "공정명을 입력해주세요."
    else:
        if not mk_number:
            return "MK번호를 입력해주세요."
        ok, msg = validate_mk(mk_number)
        if not ok:
            return msg
        if not category:
            return "공정명을 입력해주세요."
    return None


def format_duration(total_seconds):
    total_seconds = max(0, int(total_seconds))
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    seconds = total_seconds % 60
    return f"{hours}시간 {minutes}분 {seconds:02d}초", total_seconds


def finalize_work(slug, work, end_dt=None):
    end_dt = end_dt or datetime.now()
    total_seconds = int((end_dt - work["start_dt"]).total_seconds())
    duration_str, duration_seconds = format_duration(total_seconds)
    entry = {
        "id": uuid.uuid4().hex,
        "date": work["start_dt"].strftime("%Y-%m-%d"),
        "start_time": work["start_dt"].strftime("%H시 %M분"),
        "end_time": end_dt.strftime("%H시 %M분"),
        "code": work["code"],
        "mk_number": work["mk_number"],
        "category": work["category"],
        "task": work["task"],
        "duration_str": duration_str,
        "duration_seconds": duration_seconds,
        "end_timestamp": end_dt.strftime("%Y-%m-%d %H:%M:%S"),
    }
    logs = load_logs(slug)
    logs.append(entry)
    save_logs(slug, logs)
    return entry


def sort_logs(logs, reverse=False):
    return sorted(logs, key=lambda x: (x.get("date", ""), x.get("end_timestamp", "")), reverse=reverse)


# ------------------------------------------------------------------
# 페이지
# ------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ------------------------------------------------------------------
# 로그인 (이름만 입력, 비밀번호 없음)
# ------------------------------------------------------------------
@app.route("/api/me")
def api_me():
    slug = session.get("user_slug")
    return jsonify({"logged_in": bool(slug), "name": session.get("user_name")})


@app.route("/api/login", methods=["POST"])
def api_login():
    body = request.get_json(force=True, silent=True) or {}
    name = (body.get("name") or "").strip()
    if not name:
        return jsonify({"error": "이름을 입력해주세요."}), 400
    if len(name) > 20:
        return jsonify({"error": "이름은 20자 이내로 입력해주세요."}), 400
    slug = slugify_user(name)
    session["user_slug"] = slug
    session["user_name"] = name
    session.permanent = True
    user_dir(slug)
    return jsonify({"ok": True, "name": name})


@app.route("/api/logout", methods=["POST"])
def api_logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/meta")
def api_meta():
    return jsonify({"categories": CATEGORY_LIST, "category_map": CATEGORY_CODE_MAP, "codes": CODES})


# ------------------------------------------------------------------
# 진행중인 업무 (시작 / 종료 / 초기화)
# ------------------------------------------------------------------
@app.route("/api/work")
def api_work_get():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    with ACTIVE_LOCK:
        w = ACTIVE_WORK.get(slug)
        if not w:
            return jsonify({"active": None})
        elapsed = int((datetime.now() - w["start_dt"]).total_seconds())
        return jsonify({
            "active": {
                "code": w["code"],
                "mk_number": w["mk_number"],
                "category": w["category"],
                "task": w["task"],
                "start_time": w["start_dt"].strftime("%H:%M:%S"),
                "elapsed_seconds": elapsed,
            }
        })


@app.route("/api/work/start", methods=["POST"])
def api_work_start():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    code = (body.get("code") or "").strip()
    mk_number = normalize_mk((body.get("mk_number") or "").strip())
    category = (body.get("category") or "").strip()
    task = (body.get("task") or "").strip()
    force = bool(body.get("force"))

    with ACTIVE_LOCK:
        existing = ACTIVE_WORK.get(slug)
        if existing and not force:
            return jsonify({
                "needs_confirm": True,
                "message": "이미 작업이 진행 중입니다. 새로운 작업을 시작하시겠습니까?",
            }), 409

        err = validate_start_fields(code, mk_number, category)
        if err:
            return jsonify({"error": err}), 400

        closed_previous = None
        if existing:
            closed_previous = finalize_work(slug, existing)

        ACTIVE_WORK[slug] = {
            "start_dt": datetime.now(),
            "code": code,
            "mk_number": mk_number,
            "category": category,
            "task": task,
        }
        return jsonify({
            "ok": True,
            "start_time": ACTIVE_WORK[slug]["start_dt"].strftime("%H:%M:%S"),
            "closed_previous": closed_previous,
        })


@app.route("/api/work/end", methods=["POST"])
def api_work_end():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    with ACTIVE_LOCK:
        w = ACTIVE_WORK.pop(slug, None)
        if not w:
            return jsonify({"error": "진행 중인 작업이 없습니다."}), 400
        entry = finalize_work(slug, w)
        return jsonify({"ok": True, "entry": entry})


@app.route("/api/work/reset", methods=["POST"])
def api_work_reset():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    with ACTIVE_LOCK:
        ACTIVE_WORK.pop(slug, None)
    return jsonify({"ok": True})


# ------------------------------------------------------------------
# 업무 로그 (조회 / 수동추가 / 수정 / 삭제 / 엑셀 내보내기)
# ------------------------------------------------------------------
@app.route("/api/logs")
def api_logs_list():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    mode = request.args.get("mode", "today")
    reverse = request.args.get("reverse") == "1"
    logs = load_logs(slug)
    if mode == "today":
        today = datetime.now().strftime("%Y-%m-%d")
        logs = [l for l in logs if l.get("date") == today]
    logs = sort_logs(logs, reverse)
    return jsonify({"logs": logs, "mode": mode, "reverse": reverse})


def _minutes_diff_duration(sh, sm, eh, em):
    start_total = sh * 60 + sm
    end_total = eh * 60 + em
    if end_total >= start_total:
        diff = end_total - start_total
    else:
        diff = 0
    hours, minutes = divmod(diff, 60)
    return hours, minutes, 0


@app.route("/api/logs", methods=["POST"])
def api_logs_add():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    try:
        year = int(body.get("year"))
        month = int(body.get("month"))
        day = int(body.get("day"))
        sh, sm = int(body.get("start_h", 0)), int(body.get("start_m", 0))
        eh, em = int(body.get("end_h", 0)), int(body.get("end_m", 0))
        date_str = f"{year:04d}-{month:02d}-{day:02d}"
        datetime.strptime(date_str, "%Y-%m-%d")
    except Exception:
        return jsonify({"error": "날짜/시간 형식이 올바르지 않습니다."}), 400

    mk_number = normalize_mk((body.get("mk_number") or "").strip())
    if mk_number:
        ok, msg = validate_mk(mk_number)
        if not ok:
            return jsonify({"error": msg}), 400

    hours, minutes, seconds = _minutes_diff_duration(sh, sm, eh, em)
    duration_str = f"{hours}시간 {minutes}분 {seconds:02d}초"
    duration_seconds = hours * 3600 + minutes * 60 + seconds

    entry = {
        "id": uuid.uuid4().hex,
        "date": date_str,
        "start_time": f"{sh:02d}시 {sm:02d}분",
        "end_time": f"{eh:02d}시 {em:02d}분",
        "code": (body.get("code") or "").strip(),
        "mk_number": mk_number,
        "category": (body.get("category") or "").strip(),
        "task": (body.get("task") or "").strip(),
        "duration_str": duration_str,
        "duration_seconds": duration_seconds,
        "end_timestamp": f"{date_str} {eh:02d}:{em:02d}:00",
    }
    logs = load_logs(slug)
    logs.append(entry)
    save_logs(slug, logs)
    return jsonify({"ok": True, "entry": entry})


@app.route("/api/logs/<log_id>", methods=["PUT"])
def api_logs_edit(log_id):
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    try:
        year = int(body.get("year"))
        month = int(body.get("month"))
        day = int(body.get("day"))
        sh, sm = int(body.get("start_h", 0)), int(body.get("start_m", 0))
        eh, em = int(body.get("end_h", 0)), int(body.get("end_m", 0))
        dh, dm, ds = int(body.get("duration_h", 0)), int(body.get("duration_m", 0)), int(body.get("duration_s", 0))
        date_str = f"{year:04d}-{month:02d}-{day:02d}"
        datetime.strptime(date_str, "%Y-%m-%d")
    except Exception:
        return jsonify({"error": "날짜/시간 형식이 올바르지 않습니다."}), 400

    mk_number = normalize_mk((body.get("mk_number") or "").strip())
    if mk_number:
        ok, msg = validate_mk(mk_number)
        if not ok:
            return jsonify({"error": msg}), 400

    logs = load_logs(slug)
    for l in logs:
        if l["id"] == log_id:
            duration_seconds = dh * 3600 + dm * 60 + ds
            l.update({
                "date": date_str,
                "start_time": f"{sh:02d}시 {sm:02d}분",
                "end_time": f"{eh:02d}시 {em:02d}분",
                "code": (body.get("code") or "").strip(),
                "mk_number": mk_number,
                "category": (body.get("category") or "").strip(),
                "task": (body.get("task") or "").strip(),
                "duration_str": f"{dh}시간 {dm}분 {ds:02d}초",
                "duration_seconds": duration_seconds,
                "end_timestamp": f"{date_str} {eh:02d}:{em:02d}:00",
            })
            save_logs(slug, logs)
            return jsonify({"ok": True, "entry": l})
    return jsonify({"error": "해당 로그를 찾을 수 없습니다."}), 404


@app.route("/api/logs/delete", methods=["POST"])
def api_logs_delete():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    ids = set(body.get("ids") or [])
    if not ids:
        return jsonify({"error": "삭제할 로그를 선택해주세요."}), 400
    logs = load_logs(slug)
    remaining = [l for l in logs if l["id"] not in ids]
    deleted_count = len(logs) - len(remaining)
    save_logs(slug, remaining)
    return jsonify({"ok": True, "deleted": deleted_count})


@app.route("/api/logs/export")
def api_logs_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    slug = current_user_slug()
    if not slug:
        return login_required_response()
    mode = request.args.get("mode", "today")
    logs = load_logs(slug)
    if mode == "today":
        today = datetime.now().strftime("%Y-%m-%d")
        logs = [l for l in logs if l.get("date") == today]
    logs = sort_logs(logs, reverse=False)

    if not logs:
        return jsonify({"error": "내보낼 데이터가 없습니다."}), 400

    headers = ["날짜", "시작시간", "완료시간", "코드", "MK번호", "공정명", "업무 내용", "소요시간(분)", "누적시간(분)"]
    wb = Workbook()
    ws = wb.active
    ws.title = "업무기록"
    ws.append(headers)

    cumulative = 0
    for l in logs:
        minutes = int(round(l.get("duration_seconds", 0) / 60))
        cumulative += minutes
        ws.append([
            l.get("date", ""), l.get("start_time", ""), l.get("end_time", ""),
            l.get("code", ""), l.get("mk_number", ""), l.get("category", ""),
            l.get("task", ""), minutes, cumulative,
        ])

    column_widths = {"A": 15, "B": 15, "C": 15, "D": 10, "E": 15, "F": 30, "G": 30, "H": 15, "I": 15}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = header_alignment

    for row in ws.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            if i in (0, 1, 2, 3, 4, 7, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    label = "금일" if mode == "today" else "전체"
    filename = f"{label} 업무기록_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ------------------------------------------------------------------
# 즐겨찾기 (8개)
# ------------------------------------------------------------------
@app.route("/api/favorites")
def api_favorites_list():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    settings = load_settings(slug)
    return jsonify({"favorites": settings["favorites"]})


@app.route("/api/favorites/<num>", methods=["PUT"])
def api_favorites_save(num):
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    if num not in [str(n) for n in range(1, 9)]:
        return jsonify({"error": "잘못된 즐겨찾기 번호입니다."}), 400
    body = request.get_json(force=True, silent=True) or {}
    name = (body.get("name") or "").strip()[:15]
    mk_number = normalize_mk((body.get("mk_number") or "").strip())
    code = (body.get("code") or "").strip()
    category = (body.get("category") or "").strip()
    task = (body.get("task") or "").strip()

    if mk_number:
        ok, msg = validate_mk(mk_number)
        if not ok:
            return jsonify({"error": msg}), 400
    if code in ("CAD", "ET", "KT") and not mk_number:
        return jsonify({"error": "CAD/ET/KT 코드는 MK번호가 필요합니다."}), 400

    settings = load_settings(slug)
    settings["favorites"][num] = {
        "name": name,
        "mk_number": mk_number,
        "category": category,
        "code": code,
        "task": task,
    }
    save_settings(slug, settings)
    return jsonify({"ok": True, "favorite": settings["favorites"][num]})


@app.route("/api/favorites/<num>", methods=["DELETE"])
def api_favorites_delete(num):
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    settings = load_settings(slug)
    settings["favorites"].pop(num, None)
    save_settings(slug, settings)
    return jsonify({"ok": True})


# ------------------------------------------------------------------
# 해야할 일 (Todo)
# ------------------------------------------------------------------
def compute_todo_display(todo):
    parts = [p for p in (todo.get("mk_number"), todo.get("category"), todo.get("task")) if p]
    display_text = " - ".join(parts) if parts else "내용 없음"

    deadline = todo.get("deadline") or ""
    formatted_deadline = "N"
    dday_text = "N"
    if deadline:
        try:
            weekday_names = ["월", "화", "수", "목", "금", "토", "일"]
            deadline_date = datetime.strptime(deadline, "%Y-%m-%d")
            formatted_deadline = f"{deadline_date.strftime('%m-%d')} ({weekday_names[deadline_date.weekday()]})"
            d_day = (deadline_date.date() - datetime.now().date()).days
            if d_day == 0:
                dday_text = "D-day"
            elif d_day > 0:
                dday_text = f"D-{d_day}"
            else:
                dday_text = f"D+{abs(d_day)}"
        except Exception:
            formatted_deadline = deadline
            dday_text = "날짜오류"

    return display_text, formatted_deadline, dday_text


def _dday_sort_key(todo):
    deadline = todo.get("deadline") or ""
    if not deadline:
        return 9999
    try:
        d = (datetime.strptime(deadline, "%Y-%m-%d").date() - datetime.now().date()).days
        return -d
    except Exception:
        return 9999


@app.route("/api/todos")
def api_todos_list():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    sort_col = request.args.get("sort", "created_date")
    reverse = request.args.get("reverse") == "1"
    settings = load_settings(slug)
    todos = settings["todos"]

    important, normal = [], []
    for todo_id, todo in todos.items():
        display_text, formatted_deadline, dday_text = compute_todo_display(todo)
        row = {
            "id": todo_id,
            "mk_number": todo.get("mk_number", ""),
            "category": todo.get("category", ""),
            "code": todo.get("code", ""),
            "task": todo.get("task", ""),
            "deadline": todo.get("deadline", ""),
            "created_date": todo.get("created_date", ""),
            "important": bool(todo.get("important", False)),
            "display_text": display_text,
            "formatted_deadline": formatted_deadline,
            "dday_text": dday_text,
        }
        (important if row["important"] else normal).append(row)

    if sort_col in TODO_SORTABLE_COLUMNS:
        if sort_col == "created_date":
            normal.sort(key=lambda r: r["created_date"], reverse=reverse)
        elif sort_col == "deadline":
            normal.sort(key=lambda r: r["formatted_deadline"], reverse=reverse)
        elif sort_col == "dday":
            normal.sort(key=lambda r: _dday_sort_key(todos[r["id"]]), reverse=reverse)

    return jsonify({"todos": important + normal, "sort": sort_col, "reverse": reverse})


@app.route("/api/todos", methods=["POST"])
def api_todos_add():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    mk_number = normalize_mk((body.get("mk_number") or "").strip())
    code = (body.get("code") or "").strip()

    if mk_number:
        ok, msg = validate_mk(mk_number)
        if not ok:
            return jsonify({"error": msg}), 400
    if code in ("CAD", "ET", "KT") and not mk_number:
        return jsonify({"error": "CAD/ET/KT 코드는 MK번호가 필요합니다."}), 400

    deadline = (body.get("deadline") or "").strip()
    if deadline:
        try:
            datetime.strptime(deadline, "%Y-%m-%d")
        except Exception:
            return jsonify({"error": "기한 형식이 올바르지 않습니다. (예: 2026-01-31)"}), 400

    settings = load_settings(slug)
    todo_id = uuid.uuid4().hex
    settings["todos"][todo_id] = {
        "mk_number": mk_number,
        "category": (body.get("category") or "").strip(),
        "code": code,
        "task": (body.get("task") or "").strip(),
        "deadline": deadline,
        "created_date": datetime.now().strftime("%Y-%m-%d"),
        "important": False,
    }
    save_settings(slug, settings)
    return jsonify({"ok": True, "id": todo_id})


@app.route("/api/todos/<todo_id>", methods=["PUT"])
def api_todos_edit(todo_id):
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    settings = load_settings(slug)
    if todo_id not in settings["todos"]:
        return jsonify({"error": "해당 항목을 찾을 수 없습니다."}), 404

    body = request.get_json(force=True, silent=True) or {}
    mk_number = normalize_mk((body.get("mk_number") or "").strip())
    code = (body.get("code") or "").strip()

    if mk_number:
        ok, msg = validate_mk(mk_number)
        if not ok:
            return jsonify({"error": msg}), 400
    if code in ("CAD", "ET", "KT") and not mk_number:
        return jsonify({"error": "CAD/ET/KT 코드는 MK번호가 필요합니다."}), 400

    deadline = (body.get("deadline") or "").strip()
    if deadline:
        try:
            datetime.strptime(deadline, "%Y-%m-%d")
        except Exception:
            return jsonify({"error": "기한 형식이 올바르지 않습니다. (예: 2026-01-31)"}), 400

    todo = settings["todos"][todo_id]
    todo.update({
        "mk_number": mk_number,
        "category": (body.get("category") or "").strip(),
        "code": code,
        "task": (body.get("task") or "").strip(),
        "deadline": deadline,
    })
    save_settings(slug, settings)
    return jsonify({"ok": True})


@app.route("/api/todos/delete", methods=["POST"])
def api_todos_delete():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    ids = body.get("ids") or []
    settings = load_settings(slug)
    for tid in ids:
        settings["todos"].pop(tid, None)
    save_settings(slug, settings)
    return jsonify({"ok": True})


@app.route("/api/todos/toggle-important", methods=["POST"])
def api_todos_toggle_important():
    slug = current_user_slug()
    if not slug:
        return login_required_response()
    body = request.get_json(force=True, silent=True) or {}
    ids = body.get("ids") or []
    settings = load_settings(slug)
    for tid in ids:
        if tid in settings["todos"]:
            settings["todos"][tid]["important"] = not settings["todos"][tid].get("important", False)
    save_settings(slug, settings)
    return jsonify({"ok": True})


if __name__ == "__main__":
    from waitress import serve

    print("MK 업무기록 웹 서버 시작: http://192.168.134.130:8003")
    serve(app, host="0.0.0.0", port=8003)
