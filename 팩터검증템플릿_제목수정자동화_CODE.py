import tkinter as tk
from tkinter import filedialog
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
import re
import os

def open_folder(path):
    folder = os.path.dirname(path)
    try:
        os.startfile(folder)
    except Exception as e:
        print(f"❌ 폴더 열기 실패: {e}")

def get_user_inputs():
    pptx_path = filedialog.askopenfilename(
        title="수정할 PPTX 파일을 선택하세요",
        filetypes=[("PowerPoint Files", "*.pptx")]
    )
    if not pptx_path:
        print("❌ PPTX 파일을 선택하지 않았습니다.")
        exit()

    input_window = tk.Tk()
    input_window.title("MP번호 및 품명 입력")
    input_window.resizable(False, False)
    
    # 창을 화면 중앙에 위치시키고 최상위로 가져오기
    input_window.geometry("+%d+%d" % (input_window.winfo_screenwidth()//2 - 200, input_window.winfo_screenheight()//2 - 100))
    input_window.lift()
    input_window.focus_force()

    tk.Label(input_window, text="MP번호 (예: 24-123)").grid(row=0, column=0, padx=10, pady=5)
    tk.Label(input_window, text="MP").grid(row=0, column=1)

    def validate_2digit(P): return len(P) <= 2 and (P.isdigit() or P == "")
    def validate_3digit(P): return len(P) <= 3 and (P.isdigit() or P == "")

    vcmd2 = (input_window.register(validate_2digit), '%P')
    vcmd3 = (input_window.register(validate_3digit), '%P')

    mp_front = tk.Entry(input_window, width=3, validate='key', validatecommand=vcmd2)
    mp_front.grid(row=0, column=2)
    tk.Label(input_window, text="-").grid(row=0, column=3)
    mp_back = tk.Entry(input_window, width=5, validate='key', validatecommand=vcmd3)
    mp_back.grid(row=0, column=4)
    
    # MP번호 앞자리 입력란에 포커스 설정
    mp_front.focus()

    tk.Label(input_window, text="품명 (예: ME HL MAIN BEZEL)").grid(row=1, column=0, padx=10, pady=5)
    name_entry = tk.Entry(input_window, width=25)
    name_entry.grid(row=1, column=1, columnspan=4, padx=5)

    result = {}

    def submit(event=None):
        try:
            front = mp_front.get().strip()
            back = mp_back.get().strip()
            if not re.match(r"^\d{2}$", front):
                print("❌ MP 앞자리는 숫자 2자리여야 합니다.")
                return
            if not re.match(r"^\d{1,3}$", back):
                print("❌ MP 뒷자리는 숫자 1~3자리여야 합니다.")
                return
            mp_number = f" MP{front}-{back}"
            product_name = name_entry.get().strip()
            if not product_name:
                print("❌ 품명을 입력해주세요.")
                return
            result['pptx_path'] = pptx_path
            result['mp_number'] = mp_number
            result['product_name'] = product_name
            input_window.destroy()
        except Exception as e:
            print(f"❌ 입력 처리 중 오류 발생: {e}")
            return

    input_window.bind('<Return>', submit)
    tk.Button(input_window, text="확인", command=submit).grid(row=2, column=0, columnspan=5, pady=10)
    input_window.mainloop()

    return result

def replace_slide_titles(prs, mp_number, product_name):
    # 패턴: MP번호 다음 괄호든 아니든 품명 전체 감지
    pattern = re.compile(r"MP\d{2}-\d{3}(?:\s+\(.*?\)|\s+[^–\n]*)")

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if pattern.search(text):
                        rest_text = re.split(r"–", text, maxsplit=1)[-1].strip() if "–" in text else ""
                        paragraph.clear()

                        run1 = paragraph.add_run()
                        run1.text = f"{mp_number} "
                        run1.font.size = Pt(22)

                        run2 = paragraph.add_run()
                        run2.text = product_name
                        run2.font.color.rgb = RGBColor(0, 102, 204)
                        run2.font.size = Pt(22)

                        if rest_text:
                            run3 = paragraph.add_run()
                            run3.text = f" – {rest_text}"
                            run3.font.size = Pt(22)

def main():
    try:
        result = get_user_inputs()
        pptx_path = result['pptx_path']
        mp_number = result['mp_number']
        product_name = result['product_name']

        prs = Presentation(pptx_path)
        replace_slide_titles(prs, mp_number, product_name)

        # 파일명에 사용할 수 없는 특수문자 제거
        safe_product_name = re.sub(r'[<>:"/\\|?*]', '', product_name)
        default_name = f"{mp_number.strip()} {safe_product_name}_팩터검증템플릿_250000(MK-TECH).pptx"
        save_path = filedialog.asksaveasfilename(
            title="수정된 파일 저장 위치 선택",
            defaultextension=".pptx",
            initialfile=default_name,
            filetypes=[("PowerPoint Files", "*.pptx")]
        )

        if not save_path:
            print("❌ 저장 경로를 선택하지 않았습니다.")
            return

        prs.save(save_path)
        print(f"\n✅ 저장 완료: {save_path}")
        open_folder(save_path)

    except Exception as e:
        print(f"\n❌ 실패: {e}")

if __name__ == "__main__":
    main()
